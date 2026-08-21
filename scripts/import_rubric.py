from __future__ import annotations

import argparse
from pathlib import Path

import yaml
from openpyxl import load_workbook

from medical_evaluation.rubric import parse_time_range, split_criteria

JUDGE_TYPES = [
    "geometry",
    "hybrid",
    "state_machine",
    "reference_vlm",
    "hybrid",
    "state_machine",
    "hybrid",
    "hybrid",
    "geometry",
    "hybrid",
    "geometry",
]

REQUIRED_OBJECTS = [
    ["rubber_dam", "mark"],
    ["punch_disk", "punch_hole", "probe", "green_residue"],
    ["punch_tip", "rubber_dam", "punched_hole"],
    ["clamp"],
    ["rubber_dam", "clamp_wing", "clamp_jaw", "clamp_arm", "clamp_bow"],
    ["clamp_forceps", "clamp", "forceps_spring"],
    ["target_tooth", "adjacent_tooth", "clamp_jaw", "clamp_bow", "rubber_dam"],
    ["instrument", "clamp_wing", "target_tooth", "rubber_dam", "wing_hole"],
    ["rubber_dam_frame", "rubber_dam"],
    ["target_tooth", "adjacent_tooth", "dental_floss"],
    ["rubber_dam", "rubber_dam_frame", "nose_mouth_region"],
]


def import_rubric(source: Path) -> dict:
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        rows = list(worksheet.iter_rows(min_row=2, max_col=3, values_only=True))
    finally:
        workbook.close()
    if len(rows) != 11:
        raise ValueError(f"expected 11 rubric rows, found {len(rows)}")

    checkpoints = []
    for index, (time_text, name, criteria_text) in enumerate(rows, start=1):
        if not all(isinstance(value, str) and value.strip() for value in (time_text, name, criteria_text)):
            raise ValueError(f"rubric row {index + 1} contains an empty cell")
        checkpoints.append(
            {
                "id": f"cp_{index:02d}",
                "name": name.strip(),
                "criteria": split_criteria(criteria_text),
                "weight": 1.0 / 11.0,
                "judge_type": JUDGE_TYPES[index - 1],
                "required_objects": REQUIRED_OBJECTS[index - 1],
                "reference_time": parse_time_range(time_text).model_dump(),
                "thresholds": {},
            }
        )
    return {"version": "2026-08-21", "checkpoints": checkpoints}


def main() -> None:
    parser = argparse.ArgumentParser(description="Import the reviewed rubber-dam rubric")
    parser.add_argument("source", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    payload = import_rubric(args.source)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("w", encoding="utf-8", newline="\n") as handle:
        yaml.safe_dump(payload, handle, allow_unicode=True, sort_keys=False)
    print(f"Imported {len(payload['checkpoints'])} checkpoints to {args.output}")


if __name__ == "__main__":
    main()
